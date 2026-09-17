"""Download, verify and cache the Cook et al. 2019 hermaphrodite connectome.

Nothing here is invented. The only inputs are the published spreadsheet and the block headers
inside it. Every transformation applied is printed and is also written into PROVENANCE.md.

Usage:
    python scripts/fetch_connectome.py [--force]
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import urllib.request
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
CACHE_DIR = ROOT / "data" / "cache"

URL = (
    "https://wormwiring.org/si/"
    "SI%205%20Connectome%20adjacency%20matrices,%20corrected%20July%202020.xlsx"
)
RAW_NAME = "SI5_Connectome_adjacency_matrices_corrected_July_2020.xlsx"
# sha256 of the file as served by wormwiring.org, recorded 2026-09-17.
EXPECTED_SHA256 = "1f4fdbf84746b69b49a8da0816f52787860ce349b638dce37924ba80f90c70c9"

CHEM_SHEET = "hermaphrodite chemical"
GAP_SHEET = "hermaphrodite gap jn symmetric"

# Spreadsheet block header -> our neuron class. See DECISIONS.md D005.
BLOCK_TO_CLASS = {
    "PHARYNX": "pharyngeal",
    "SENSORY NEURONS": "sensory",
    "INTERNEURONS": "inter",
    "MOTOR NEURONS": "motor",
    "SEX SPECIFIC": "motor",
    "SEX-SPECIFIC CELLS": "motor",
}
# Neurons present in the gap sheet but absent from the chemical sheet. They are real neurons of the
# 302; the gap sheet files them under a block header that does not describe them (DECISIONS.md D004).
EXTRA_NEURONS = {"CANL": "other", "CANR": "other"}

LABEL_COL = 2  # column index holding the row's cell name
BLOCK_COL = 0  # column index holding the block header, set only on the first row of a block
NAME_ROW = 2  # row index holding the column cell names
BLOCK_ROW = 0  # row index holding the column block headers


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def download(force: bool) -> Path:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    path = RAW_DIR / RAW_NAME
    if path.exists() and not force:
        print(f"already downloaded: {path}")
    else:
        print(f"downloading {URL}")
        with urllib.request.urlopen(URL, timeout=180) as resp:
            path.write_bytes(resp.read())
        print(f"wrote {path} ({path.stat().st_size} bytes)")
    digest = sha256_of(path)
    print(f"sha256 : {digest}")
    if digest != EXPECTED_SHA256:
        print(f"EXPECTED: {EXPECTED_SHA256}")
        raise SystemExit(
            "sha256 mismatch: the upstream file changed. Do not proceed automatically -- inspect "
            "the new file, then update EXPECTED_SHA256 and PROVENANCE.md deliberately."
        )
    return path


def sheet_grid(workbook, sheet_name: str) -> list[tuple]:
    return list(workbook[sheet_name].iter_rows(values_only=True))


def axis_labels(grid: list[tuple], axis: str) -> tuple[dict[str, int], dict[str, str]]:
    """Return {cell name: index} and {cell name: block header} for one axis of a sheet."""
    index: dict[str, int] = {}
    block: dict[str, str] = {}
    current = None
    if axis == "row":
        headers = {i: row[BLOCK_COL] for i, row in enumerate(grid) if row[BLOCK_COL]}
        for i, row in enumerate(grid):
            if i in headers:
                current = headers[i]
            name = row[LABEL_COL] if len(row) > LABEL_COL else None
            if name:
                if name in index:
                    raise ValueError(f"duplicate row label {name!r}")
                index[name], block[name] = i, current
    else:
        headers = {i: v for i, v in enumerate(grid[BLOCK_ROW]) if v}
        for i, name in enumerate(grid[NAME_ROW]):
            if i in headers:
                current = headers[i]
            if name:
                if name in index:
                    raise ValueError(f"duplicate column label {name!r}")
                index[name], block[name] = i, current
    return index, block


def extract(grid: list[tuple], names: list[str]) -> np.ndarray:
    """Pull the len(names) x len(names) submatrix for `names`, zero where a label is absent."""
    rows, _ = axis_labels(grid, "row")
    cols, _ = axis_labels(grid, "col")
    out = np.zeros((len(names), len(names)), dtype=np.float32)
    for i, src in enumerate(names):
        r = rows.get(src)
        if r is None:
            continue
        row = grid[r]
        for j, dst in enumerate(names):
            c = cols.get(dst)
            if c is None or c >= len(row):
                continue
            v = row[c]
            if v is None or v == "":
                continue
            out[i, j] = float(v)
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="re-download even if the file is present")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("openpyxl is required: pip install openpyxl==3.1.5")
        return 1

    xlsx = download(args.force)
    print(f"reading {xlsx.name}")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    for sheet in (CHEM_SHEET, GAP_SHEET):
        if sheet not in wb.sheetnames:
            raise SystemExit(f"sheet {sheet!r} missing; sheets are {wb.sheetnames}")

    chem_grid = sheet_grid(wb, CHEM_SHEET)
    gap_grid = sheet_grid(wb, GAP_SHEET)

    chem_rows, chem_row_block = axis_labels(chem_grid, "row")
    gap_rows, _ = axis_labels(gap_grid, "row")

    # Neuron list and ordering: every cell that is a row of the chemical sheet (all of them are
    # neurons; muscles and end organs appear only as columns there), in sheet order, then the
    # neurons that exist only in the gap sheet.
    names: list[str] = [n for n in chem_rows if chem_row_block[n] in BLOCK_TO_CLASS]
    unknown = [n for n in chem_rows if chem_row_block[n] not in BLOCK_TO_CLASS]
    if unknown:
        raise SystemExit(f"unmapped chemical-sheet blocks for {unknown[:5]} -- check BLOCK_TO_CLASS")
    classes: list[str] = [BLOCK_TO_CLASS[chem_row_block[n]] for n in names]
    for extra, cls in EXTRA_NEURONS.items():
        if extra not in gap_rows:
            raise SystemExit(f"{extra} expected in {GAP_SHEET!r} but not found")
        if extra in chem_rows:
            raise SystemExit(f"{extra} unexpectedly present in {CHEM_SHEET!r}; revisit D004")
        names.append(extra)
        classes.append(cls)

    print(f"neurons: {len(names)}")
    if len(names) != 302:
        raise SystemExit(f"expected 302 neurons, got {len(names)}")

    chem = extract(chem_grid, names)
    gap = extract(gap_grid, names)

    # --- transformations, each one reported ---
    notes: list[str] = []

    # Chemical autapses are in the published matrix and are kept: W_ii * tanh(v_i) is a real
    # dynamical term, and SH/RD graphs are allowed self-loops so the conditions stay matched.
    # Gap self-edges are dropped: G_ii * (v_i - v_i) is identically zero, so a nonzero G_ii would
    # only corrupt the implicit denominator of the integrator. See DECISIONS.md D007.
    self_chem = int((np.diag(chem) != 0).sum())
    self_gap = int((np.diag(gap) != 0).sum())
    np.fill_diagonal(gap, 0.0)
    notes.append(
        f"chemical autapses kept: {self_chem} (sum {float(np.diag(chem).sum()):g}); "
        f"gap self-edges dropped: {self_gap}"
    )

    asym = float(np.abs(gap - gap.T).max())
    notes.append(f"max |G - G^T| before any symmetrisation: {asym:g}")
    if asym != 0.0:
        gap = 0.5 * (gap + gap.T)
        notes.append("gap matrix symmetrised as (G + G^T)/2")

    neg_chem = int((chem < 0).sum())
    neg_gap = int((gap < 0).sum())
    notes.append(f"negative entries: {neg_chem} chemical, {neg_gap} gap (none expected)")

    chem_edges = int((chem > 0).sum())
    gap_edges = int((gap > 0).sum())
    no_out = [n for i, n in enumerate(names) if chem[i].sum() == 0]
    no_in = [n for i, n in enumerate(names) if chem[:, i].sum() == 0]
    no_gap = [n for i, n in enumerate(names) if gap[i].sum() == 0]
    isolated = [n for i, n in enumerate(names) if chem[i].sum() == chem[:, i].sum() == gap[i].sum() == 0]
    notes.append(
        f"edges: {chem_edges} chemical (directed), {gap_edges} gap (directed count, "
        f"{gap_edges // 2} undirected)"
    )
    notes.append(f"neurons with no outgoing chemical synapse ({len(no_out)}): {sorted(no_out)}")
    notes.append(f"neurons with no incoming chemical synapse ({len(no_in)}): {sorted(no_in)}")
    notes.append(f"neurons with no gap junction ({len(no_gap)}): {sorted(no_gap)}")
    notes.append(
        f"fully isolated neurons ({len(isolated)}): {sorted(isolated)} "
        "-- kept in the 302 regardless, so the node set is identical across N2/SH/RD"
    )

    meta = {
        "source_url": URL,
        "raw_file": RAW_NAME,
        "sha256": EXPECTED_SHA256,
        "chem_sheet": CHEM_SHEET,
        "gap_sheet": GAP_SHEET,
        "weight_kind": "em_sections",
        "weight_meaning": (
            "total number of EM serial sections of connectivity, taking into account both the "
            "number of synapses and the sizes of synapses (Cook et al. 2019, SI 5 legend). "
            "This is NOT a synapse count."
        ),
        "citation": (
            "Cook SJ, Jarrell TA, Brittin CA, Wang Y, Bloniarz AE, Yakovlev MA, Nguyen KCQ, "
            "Tang LT-H, Bayer EA, Duerr JS, Bulow HE, Hobert O, Hall DH, Emmons SW (2019). "
            "Whole-animal connectomes of both Caenorhabditis elegans sexes. Nature 571:63-71."
        ),
        "n_neurons": len(names),
        "chem_edges": chem_edges,
        "gap_edges_directed": gap_edges,
        "transformations": notes,
    }

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    out = CACHE_DIR / "cook2019_herm.npz"
    np.savez_compressed(
        out,
        names=np.array(names),
        classes=np.array(classes),
        chem=chem,
        gap=gap,
        meta=np.array(json.dumps(meta, indent=2)),
    )
    print(f"\nwrote {out}")
    for note in notes:
        print(f"  - {note}")
    from collections import Counter

    print(f"  - classes: {dict(Counter(classes))}")
    print(f"  - cache sha256: {sha256_of(out)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
